#!/usr/bin/env python3
"""Genere le guide du Coach JuiceLab pour PwnzzAI au format .docx.

Modele : docs/generate_install_guides.py du depot JuiceLab (memes helpers
python-docx, meme palette, diagrammes Mermaid rendus en PNG via mmdc).

Produit un document Word distribuable (prof + eleve) :

    coach/docs/GUIDE-COACH-PWNZZAI.docx

Dependances :
    - python-docx               (pip install python-docx)
    - @mermaid-js/mermaid-cli    -> binaire `mmdc` (npm i -g @mermaid-js/mermaid-cli)
      + un Chromium/Chrome (mmdc en a besoin). Sur systeme sans le chromium
      bundle de puppeteer, exporter PUPPETEER_EXECUTABLE_PATH vers le binaire
      systeme (ex. /usr/bin/chromium-browser).

Usage :
    python coach/docs/build_coach_guide.py

Source de verite : valeurs ci-dessous, alignees sur docker-compose.coach.yml
et le .env de PwnzzAI. Si l'une change, mettre a jour ici et relancer.
"""

from __future__ import annotations

import os
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path

try:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Inches, Mm, Pt, RGBColor
except ImportError:
    sys.exit(
        "python-docx introuvable. Installer : pip install python-docx "
        "(ou venv : python3 -m venv .venv && .venv/bin/pip install python-docx)."
    )

# ---------------------------------------------------------------------------
# Constantes projet (alignees sur docker-compose.coach.yml + .env)
# ---------------------------------------------------------------------------

COACH_PORT = "8095"
RAW_PORT = "8090"
DASHBOARD_PORT = "5000"
COHORT = "M2-IA-2026"
LABEL = "pwnzzai-poste-01"
JUDGE_MODEL = "llama3.2:3b"

DOCS_DIR = Path(__file__).resolve().parent
IMG_DIR = DOCS_DIR / "img"
OUT = DOCS_DIR / "GUIDE-COACH-PWNZZAI.docx"

# Palette (identique au modele JuiceLab, teinte violette du coach pour le titre)
BLUE = RGBColor(0x2E, 0x75, 0xB6)
PURPLE = RGBColor(0x6D, 0x28, 0xD9)
DARK = RGBColor(0x22, 0x22, 0x22)
GREY_FILL = "F2F2F2"
NOTE_FILL = "FFF4D6"
HEAD_FILL = "EDE7FB"

# ---------------------------------------------------------------------------
# Rendu Mermaid -> PNG
# ---------------------------------------------------------------------------


def _find_mmdc() -> str:
    """Prefere le mmdc local (coach/docs/node_modules) au mmdc global.

    Sur les distros a snap (Ubuntu), le Chromium snap ne peut pas lire les
    fichiers de mmdc s'il est installe sous un dossier cache comme ~/.nvm.
    Un mmdc installe dans ce dossier (chemin non cache) contourne le souci.
    """
    local = DOCS_DIR / "node_modules" / ".bin" / "mmdc"
    if local.is_file():
        return str(local)
    return shutil.which("mmdc") or ""


def _find_chromium() -> str:
    """Chromium utilisable pour mmdc. PUPPETEER_EXECUTABLE_PATH gagne, sinon
    on tente les binaires systeme courants (le Chromium snap fonctionne)."""
    exe = os.environ.get("PUPPETEER_EXECUTABLE_PATH", "")
    if exe and Path(exe).exists():
        return exe
    for cand in ("/snap/bin/chromium", "/usr/bin/chromium-browser",
                 "/usr/bin/chromium", "/usr/bin/google-chrome"):
        if Path(cand).exists():
            return cand
    return ""


def render_mermaid(source: str, out_png: Path, tmp: Path) -> bool:
    """Rend un diagramme Mermaid en PNG via mmdc. Retourne True si succes."""
    mmdc = _find_mmdc()
    if not mmdc:
        print("  [warn] mmdc introuvable — diagramme saute. Installer : "
              "cd coach/docs && npm install @mermaid-js/mermaid-cli",
              file=sys.stderr)
        return False

    mmd = tmp / (out_png.stem + ".mmd")
    mmd.write_text(source.strip() + "\n", encoding="utf-8")

    pcfg = tmp / "puppeteer.json"
    chromium = _find_chromium()
    if chromium:
        pcfg.write_text(
            '{"executablePath":"%s","args":["--no-sandbox","--disable-gpu"]}'
            % chromium, encoding="utf-8",
        )
    else:
        pcfg.write_text('{"args":["--no-sandbox","--disable-gpu"]}',
                        encoding="utf-8")

    cmd = [mmdc, "-i", str(mmd), "-o", str(out_png),
           "-b", "white", "-s", "2", "-p", str(pcfg)]
    proc = subprocess.run(cmd, capture_output=True, text=True)
    if proc.returncode != 0 or not out_png.exists():
        print(f"  [warn] echec rendu {out_png.name} :\n{proc.stderr.strip()}",
              file=sys.stderr)
        return False
    return True


# ---------------------------------------------------------------------------
# Helpers python-docx (repris du modele JuiceLab)
# ---------------------------------------------------------------------------


def setup_styles(doc: Document) -> None:
    section = doc.sections[0]
    section.page_width = Mm(210)
    section.page_height = Mm(297)
    section.top_margin = Mm(20)
    section.bottom_margin = Mm(20)
    section.left_margin = Mm(20)
    section.right_margin = Mm(20)

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)

    for level, size in ((1, 18), (2, 14), (3, 12)):
        st = doc.styles[f"Heading {level}"]
        st.font.name = "Calibri"
        st.font.size = Pt(size)
        st.font.bold = True
        st.font.color.rgb = PURPLE if level == 1 else DARK

    zoom = doc.settings.element.find(qn("w:zoom"))
    if zoom is not None:
        zoom.set(qn("w:percent"), "100")
        if zoom.get(qn("w:val")) is not None:
            del zoom.attrib[qn("w:val")]


_PPR_AFTER_SHD = (
    "w:tabs", "w:suppressAutoHyphens", "w:kinsoku", "w:wordWrap",
    "w:overflowPunct", "w:topLinePunct", "w:autoSpaceDE", "w:autoSpaceDN",
    "w:bidi", "w:adjustRightInd", "w:snapToGrid", "w:spacing", "w:ind",
    "w:contextualSpacing", "w:mirrorIndents", "w:suppressOverlap", "w:jc",
    "w:textDirection", "w:textAlignment", "w:textboxTightWrap", "w:outlineLvl",
    "w:divId", "w:cnfStyle", "w:rPr", "w:sectPr", "w:pPrChange",
)


def _shade(paragraph, fill: str) -> None:
    pPr = paragraph._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:fill"), fill)
    pPr.insert_element_before(shd, *_PPR_AFTER_SHD)


def title_block(doc: Document, title: str, subtitle: str) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = p.add_run(title)
    run.font.size = Pt(26)
    run.font.bold = True
    run.font.color.rgb = PURPLE
    sub = doc.add_paragraph()
    r = sub.add_run(subtitle)
    r.font.size = Pt(12)
    r.font.color.rgb = DARK
    r.italic = True
    bar = doc.add_paragraph()
    pPr = bar._p.get_or_add_pPr()
    pbdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "12")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), "6D28D9")
    pbdr.append(bottom)
    pPr.append(pbdr)


def para(doc: Document, text: str) -> None:
    doc.add_paragraph(text)


def bullet(doc: Document, text: str) -> None:
    doc.add_paragraph(text, style="List Bullet")


def numbered(doc: Document, text: str) -> None:
    doc.add_paragraph(text, style="List Number")


def code_block(doc: Document, lines: list[str]) -> None:
    for ln in lines:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.space_before = Pt(0)
        _shade(p, GREY_FILL)
        run = p.add_run(ln if ln else " ")
        run.font.name = "Consolas"
        run.font.size = Pt(9.5)
    doc.add_paragraph().paragraph_format.space_after = Pt(4)


def note(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    _shade(p, NOTE_FILL)
    run = p.add_run(text)
    run.font.size = Pt(10.5)


def add_table(doc: Document, headers: list[str], rows: list[list[str]],
              widths_mm: list[int]) -> None:
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    hdr = table.rows[0].cells
    for i, h in enumerate(headers):
        hdr[i].text = ""
        run = hdr[i].paragraphs[0].add_run(h)
        run.bold = True
        run.font.size = Pt(10)
        _shade(hdr[i].paragraphs[0], HEAD_FILL)
    for row in rows:
        cells = table.add_row().cells
        for i, val in enumerate(row):
            cells[i].text = ""
            run = cells[i].paragraphs[0].add_run(val)
            run.font.size = Pt(10)
    for col_idx, w in enumerate(widths_mm):
        for cell in table.columns[col_idx].cells:
            cell.width = Mm(w)


def add_diagram(doc: Document, png: Path, caption: str) -> None:
    if png and png.exists():
        doc.add_picture(str(png), width=Inches(6.2))
        doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
        cap = doc.add_paragraph()
        cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = cap.add_run(caption)
        r.italic = True
        r.font.size = Pt(9)
        r.font.color.rgb = DARK


def add_screenshot(doc: Document, png: Path, caption: str) -> None:
    """Insere une capture d'ecran (PNG deja sur disque) + legende centree.

    Contrairement a add_diagram (PNG rendu a la volee par mmdc), la capture
    est versionnee dans docs/img/. Si absente, on signale au lieu de planter.
    """
    if png and png.exists():
        add_diagram(doc, png, caption)
    else:
        print("  [warn] capture introuvable, saut : %s" % png, file=sys.stderr)


# ---------------------------------------------------------------------------
# Diagrammes Mermaid
# ---------------------------------------------------------------------------

MERMAID_ARCHI = """
flowchart LR
    B["Navigateur eleve<br/>:%COACH%"]
    C["pwnzzai-coach<br/>proxy + API"]
    P["pwnzzai-app<br/>:8080 (OWASP intact)"]
    O["ollama"]
    D["Dashboard prof<br/>:%DASH%"]
    B -->|"toutes les requetes"| C
    C -->|"transmises telles quelles"| P
    C -->|"juge + indices"| O
    C -->|"evenements cohorte"| D
    C -. "injecte coach.js" .-> B
""".replace("%COACH%", COACH_PORT).replace("%DASH%", DASHBOARD_PORT)

MERMAID_SEQUENCE = """
sequenceDiagram
    autonumber
    participant E as Eleve
    participant A as Assistant vulnerable
    participant C as coach
    participant O as Ollama (juge)
    participant D as Dashboard prof
    E->>A: prompts d'attaque
    A-->>E: reponses
    Note over E: coach.js capture la conversation
    E->>C: POST /__coach/judge {transcript}
    C->>O: prompt + few-shot + transcript
    O-->>C: VERDICT / SCORE / REASON
    C-->>E: {success, score, reason}
    E->>C: POST /__coach/event (si reussi)
    C->>D: POST /api/sync
"""


# ---------------------------------------------------------------------------
# Document
# ---------------------------------------------------------------------------


def build(tmp: Path) -> None:
    doc = Document()
    setup_styles(doc)
    title_block(
        doc, "Coach JuiceLab pour PwnzzAI — Guide",
        "OWASP PwnzzAI + couche pedagogique (juge LLM, indices, suivi cohorte)",
    )

    doc.add_heading("1. Ce que le coach apporte", level=1)
    para(doc, "Le Coach JuiceLab transforme OWASP PwnzzAI (la pizza shop "
              "volontairement vulnerable pour apprendre la securite des LLM) en "
              "un TD guide et suivi a distance, sans modifier une seule ligne du "
              "produit OWASP.")
    para(doc, "Il apporte la parite de fonctionnalites eleve de JuiceLab via une "
              "sidebar a onglets injectee dans le navigateur :")
    bullet(doc, "Briefing : mission + concepts pedagogiques par lab (bilingue), "
                "ancres OWASP LLM Top 10.")
    bullet(doc, "Indices gradues : 5 niveaux N1-N5, cout 5/10/20/35/50 %, "
                "revelation progressive, penalite de score. Contenu genere par "
                "le modele local, adapte aux tentatives ratees.")
    bullet(doc, "Journal avant/apres et Quiz (3 QCM corriges avec explications).")
    bullet(doc, "Juge automatique (LLM-as-judge) : verdict + score 0-100 + "
                "justification, a partir de la conversation capturee.")
    bullet(doc, "Onglet Progression (dashboard eleve) : score par lab, indices "
                "consommes, badges (4 tiers), score moyen.")
    bullet(doc, "Remontee vers le dashboard prof JuiceLab : un TD PwnzzAI "
                "apparait dans la meme matrice de cohorte que Juice Shop.")
    note(doc, "Scoring (identique a JuiceLab) : score = max(50, 100 - somme des "
              "couts d'indices).")

    png_archi = tmp / "archi.png"
    render_mermaid(MERMAID_ARCHI, png_archi, tmp)
    add_diagram(doc, png_archi, "Le coach est un sidecar place devant PwnzzAI ; "
                                "l'image OWASP reste intacte.")

    doc.add_heading("2. Ce qui s'installe", level=1)
    add_table(
        doc, ["Conteneur", "Port", "Role"],
        [
            ["pwnzzai-coach", COACH_PORT, "Entree eleve : proxy + panneau coach"],
            ["pwnzzai-shop", RAW_PORT, "PwnzzAI brut OWASP (intact, debug)"],
            ["ollama", "interne", "Modele des labs + modele du juge/indices"],
        ],
        [45, 25, 100],
    )
    note(doc, "L'eleve utilise uniquement http://localhost:" + COACH_PORT +
              ". Le port " + RAW_PORT + " reste disponible pour deboguer "
              "PwnzzAI sans le coach.")

    doc.add_heading("3. Installation", level=1)
    para(doc, "Stack autonome (PwnzzAI clone au build). Depuis la racine du depot :")
    code_block(doc, [
        "cp .env.example .env    # cohorte + dashboard",
        "docker compose up -d --build",
    ])
    para(doc, "Puis tirer les modeles Ollama (page Basics de PwnzzAI, ou CLI) :")
    code_block(doc, [
        "docker exec ollama ollama pull llama3.2:1b   # modele des labs",
        "docker exec ollama ollama pull " + JUDGE_MODEL + "   # modele du juge",
    ])
    para(doc, "Verifier que tout repond :")
    code_block(doc, [
        "curl -s http://localhost:" + COACH_PORT + "/__coach/health",
        '# {"ok":true,"ollama":true,"dashboard_configured":true}',
    ])

    doc.add_heading("4. Connexion au dashboard prof (la cohorte)", level=1)
    para(doc, "Topologie : un dashboard central, PwnzzAI est un client. Le "
              "dashboard prof est une instance unique et partagee, deployee une "
              "seule fois cote juicelab. Elle sert a la fois les cohortes "
              "JuiceLab (Juice Shop) ET les cohortes PwnzzAI : les deux remontent "
              "dans la meme matrice.")
    para(doc, "PwnzzAI ne duplique JAMAIS le code serveur du dashboard. Il s'y "
              "branche comme client via la variable JUICELAB_DASHBOARD_URL. "
              "Anti-pattern a proscrire : deployer un second dashboard "
              "specifique a PwnzzAI.")
    para(doc, "Cas optionnel : si le dashboard central n'existe pas encore et que "
              "tu veux le deployer depuis cette machine, le script "
              "scripts/deploy-dashboard.sh fait un sparse-checkout pinne du depot "
              "juicelab (REF reglee par JUICELAB_DASHBOARD_REF, SHA conseille en "
              "prod) et ne tire QUE la partie prof (dashboard/, docker/, "
              "scripts/). L'overlay et le Juice Shop eleve ne sont jamais tires. "
              "Une fois le dashboard demarre, pointer JUICELAB_DASHBOARD_URL "
              "dessus.")
    add_screenshot(
        doc, IMG_DIR / "prof-dashboard-light.png",
        "Tableau de bord prof - cohorte PwnzzAI sur l'instance centrale "
        "(meme matrice que les cohortes Juice Shop).")
    para(doc, "Tout se regle dans le fichier .env de PwnzzAI :")
    add_table(
        doc, ["Variable", "Role", "Exemple"],
        [
            ["JUICELAB_DASHBOARD_URL", "URL du dashboard vue depuis le conteneur",
             "http://host.docker.internal:" + DASHBOARD_PORT],
            ["JUICELAB_COHORT_ID", "Identifiant de la promo", COHORT],
            ["JUICELAB_INSTANCE_LABEL", "Nom du poste eleve dans la matrice", LABEL],
            ["COACH_JUDGE_MODEL", "Modele Ollama du juge/indices", JUDGE_MODEL],
        ],
        [60, 75, 55],
    )
    note(doc, "Dashboard sur la meme machine : utiliser host.docker.internal "
              "(et non 127.0.0.1). Dashboard distant : mettre l'IP du prof. "
              "Vide = mode local sans remontee. Le dashboard enregistre l'eleve "
              "automatiquement au premier evenement.")
    para(doc, "Apres modification du .env, recreer le conteneur coach :")
    code_block(doc, [
        "docker compose up -d pwnzzai-coach",
    ])

    doc.add_heading("5. Cote eleve : comment ca s'utilise", level=1)
    numbered(doc, "Ouvrir http://localhost:" + COACH_PORT +
                  " et aller sur un lab (ex. Direct Prompt Injection).")
    numbered(doc, "Cliquer le bouton rond violet JL en bas a droite : le "
                  "panneau coach a onglets s'ouvre (bouton FR/EN).")
    numbered(doc, "5 onglets : Briefing (mission + concepts), Indices (5 niveaux "
                  "N1-N5 a cout croissant), Journal (avant/apres), Quiz (3 QCM), "
                  "Progression (juge, score, badges).")
    numbered(doc, "Attaquer l'assistant comme d'habitude : la conversation est "
                  "capturee automatiquement.")
    numbered(doc, "Bloque ? L'onglet Indices revele un conseil gradue (N1 = "
                  "declic, N5 = exemple quasi complet) ; chaque indice baisse le "
                  "score (plancher 50).")
    numbered(doc, "Verifier ma reussite (onglet Progression) soumet la "
                  "conversation au juge : Reussi / Partiel / Pas encore + score + "
                  "justification ; une reussite peut debloquer un badge.")
    note(doc, "Astuce prof : un lien termine par #coach (ex. "
              "http://localhost:" + COACH_PORT + "/indirect-prompt-injection"
              "#coach) ouvre le panneau automatiquement.")

    doc.add_heading("6. Cote prof : ce qui remonte", level=1)
    para(doc, "Le coach envoie ces evenements via le contrat public POST "
              "/api/sync (le meme que Juice Shop) :")
    add_table(
        doc, ["Evenement", "Declencheur", "Donnees"],
        [
            ["session_start", "ouverture d'une page lab", "chemin, identite si detectee"],
            ["hint_revealed", "revelation d'un indice", "niveau N1-N5, cout, score"],
            ["challenge_solved", "juge -> Reussi", "score, verdict, justification"],
            ["journal_filled", "Journal apres enregistre", "longueur, texte"],
            ["quiz_completed", "validation du quiz", "score, bonnes / total"],
            ["badge_earned", "badge debloque", "identifiant du badge"],
        ],
        [45, 60, 65],
    )

    png_seq = tmp / "sequence.png"
    render_mermaid(MERMAID_SEQUENCE, png_seq, tmp)
    add_diagram(doc, png_seq, "Du transcript au verdict du juge, puis remontee "
                              "au dashboard si reussi.")

    doc.add_heading("7. Choisir le modele du juge", level=1)
    add_table(
        doc, ["Modele", "Verdict", "Recommandation"],
        [
            ["llama3.2:1b", "non fiable (incoherent)", "a eviter pour le juge"],
            ["llama3.2:3b", "fiable sur cas nets", "minimum conseille"],
            ["mistral:7b / llama3.1:8b", "fiable sur cas subtils",
             "ideal si la machine suit"],
        ],
        [50, 60, 60],
    )
    note(doc, "Le modele du juge est independant de celui des labs : labs en 1b "
              "(rapide) et juge en 3b (fiable) sur la meme machine. Regle par "
              "COACH_JUDGE_MODEL.")

    doc.add_heading("8. Depannage", level=1)
    add_table(
        doc, ["Symptome", "Cause", "Solution"],
        [
            ["Service coach indisponible", "aucun modele de juge tire",
             "ollama pull " + JUDGE_MODEL],
            ["health: ollama false", "idem", "tirer un modele, verifier COACH_JUDGE_MODEL"],
            ["Verdicts incoherents", "modele trop petit (1b)", "passer a 3b ou plus"],
            ["Dashboard non configure", "URL vide", "renseigner l'URL, recreer le conteneur"],
            ["Rien ne remonte", "dashboard injoignable du conteneur",
             "host.docker.internal, port/pare-feu"],
            ["Port " + COACH_PORT + " pris", "autre service",
             "changer le mapping dans docker-compose.coach.yml"],
        ],
        [50, 55, 65],
    )

    doc.add_heading("9. Pourquoi ca resiste aux evolutions d'OWASP", level=1)
    para(doc, "Le proxy ne connait aucune route interne de PwnzzAI : il transmet "
              "tout et n'injecte qu'une balise script. La capture est generique "
              "(tout POST fetch/XHR portant un champ texte). Si OWASP renomme une "
              "route ou change le style d'une page, le coach continue de "
              "fonctionner ; seul labs.json peut demander une mise a jour de "
              "chemin pour la detection du lab.")

    doc.save(str(OUT))
    print(f"  [ok] {OUT.name} genere ({OUT.stat().st_size // 1024} Ko)")


def main() -> None:
    print("Generation du guide Coach JuiceLab pour PwnzzAI...")
    # tmp sous DOCS_DIR (chemin non cache) : le Chromium snap, confine, ne lit
    # pas les dossiers caches type ~/.nvm ou /tmp restreint.
    with tempfile.TemporaryDirectory(dir=str(DOCS_DIR)) as td:
        build(Path(td))


if __name__ == "__main__":
    main()
