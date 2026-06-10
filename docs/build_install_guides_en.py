#!/usr/bin/env python3.11
"""Generate the EN teacher install guide as .docx.

English counterpart of docs/build_install_guides.py (which produces the FR
GUIDE-INSTALL-PROF.docx). This wrapper reuses ALL the low-level helpers and
constants from the FR module (styles, tables, code blocks, mermaid rendering)
and only re-states the prose in English.

    python3.11 docs/build_install_guides_en.py    # -> GUIDE-INSTALL-PROF-EN.docx
"""
from pathlib import Path
import tempfile

import build_install_guides as p
from docx import Document

DOCS_DIR = p.DOCS_DIR
PROF_OUT_EN = DOCS_DIR / "GUIDE-INSTALL-PROF-EN.docx"

MERMAID_ARCHI_EN = """
flowchart LR
    B["Student browser<br/>:%COACH%"]
    C["pwnzzai-coach<br/>proxy + API"]
    P["pwnzzai-app<br/>:8080 (OWASP intact)"]
    O["ollama"]
    D["Teacher dashboard<br/>:%DASH%"]
    B -->|"all requests"| C
    C -->|"forwarded as-is"| P
    C -->|"judge + hints"| O
    C -->|"cohort events"| D
    C -. "injects coach.js" .-> B
""".replace("%COACH%", p.COACH_PORT).replace("%DASH%", p.DASHBOARD_PORT)

MERMAID_SEQUENCE_EN = """
sequenceDiagram
    autonumber
    participant E as Student
    participant A as Vulnerable assistant
    participant C as coach
    participant O as Ollama (judge)
    participant D as Teacher dashboard
    E->>A: attack prompts
    A-->>E: replies
    Note over E: coach.js captures the conversation
    E->>C: POST /__coach/judge {transcript}
    C->>O: prompt + few-shot + transcript
    O-->>C: VERDICT / SCORE / REASON
    C-->>E: {success, score, reason}
    E->>C: POST /__coach/event (if solved)
    C->>D: POST /api/sync
"""


def section_prerequisites_en(doc: Document) -> None:
    p.add_table(
        doc, ["Tool", "Minimum version", "Note"],
        [
            ["Docker Desktop (Win/macOS) or Docker Engine (Linux)", "24+", "docker.com"],
            ["Docker Compose v2", "ships with Docker Desktop", "Linux: see Docker appendix"],
            ["Git", "recent version", "git-scm.com"],
            ["RAM", "8 GB free recommended", "Ollama loads the judge model in memory"],
            ["Disk", "6 GB free", "PwnzzAI image + Ollama models"],
        ],
        [70, 45, 55],
    )
    p.para(doc, "Quick sanity check:")
    p.code_block(doc, [
        "docker --version            # 24.x or newer",
        "docker compose version      # v2.x or newer",
        "git --version",
    ])
    p.note(doc, "If 'docker compose version' fails, Docker is too old. "
                "Linux: sudo apt install docker-compose-v2 (or docker-compose-plugin). "
                "Windows / macOS: update Docker Desktop.")


def section_llm_provider_en(doc: Document, full: bool = False) -> None:
    p.para(doc, "By default the labs' TARGET assistant runs on local Ollama "
                "(100% offline). Since the OWASP product is built on LiteLLM, this "
                "target can run on many providers (Groq, OpenAI, Gemini, Anthropic) "
                "WITHOUT any code change — pure .env config. MODEL_PROVIDER=openai "
                "means 'cloud via LiteLLM' (not necessarily OpenAI); the prefix "
                "before the / in LITELLM_MODEL picks the real provider.")
    p.add_table(
        doc, ["Provider", "MODEL_PROVIDER", "LITELLM_MODEL (example)", "Key var"],
        [
            ["Ollama (local, default)", "ollama", "— (uses OLLAMA_MODEL)", "—"],
            ["Groq (recommended)", "openai", "groq/llama-3.3-70b-versatile", "GROQ_API_KEY"],
            ["OpenAI", "openai", "openai/gpt-4o-mini", "OPENAI_API_KEY"],
            ["Google Gemini", "openai", "gemini/gemini-2.0-flash", "GEMINI_API_KEY"],
            ["Anthropic", "openai", "anthropic/claude-3-5-haiku-latest", "ANTHROPIC_API_KEY"],
        ],
        [38, 30, 60, 42],
    )
    p.para(doc, "To switch the target to a cloud provider (Groq example), put these "
                "three lines in .env then restart:")
    p.code_block(doc, [
        "MODEL_PROVIDER=openai",
        "LITELLM_MODEL=groq/llama-3.3-70b-versatile",
        "GROQ_API_KEY=gsk_xxx",
        "",
        "./pwnzzai.sh restart",
    ])
    p.note(doc, "Groq model: use groq/llama-3.3-70b-versatile (non-reasoning, "
                "reliable). Avoid groq/openai/gpt-oss-20b — a reasoning model that "
                "often returns empty (tested: HTTP 500 / empty response on some labs). "
                "Valid key = starts with gsk_ and returns 200 on "
                "https://api.groq.com/openai/v1/models (401 = invalid).")
    p.note(doc, "Only the openai_* cloud labs use this backend; ollama_* labs always "
                "use local Ollama, so keep the ollama service running. The key stays "
                "in .env (gitignored) — never commit it.")
    if full:
        p.note(doc, "Security: the compose file passes GROQ_API_KEY, OPENAI_API_KEY, "
                    "GEMINI_API_KEY and ANTHROPIC_API_KEY to pwnzzai-app, all empty by "
                    "default (no behaviour change if unset). The key lives only in "
                    ".env, never committed.")
        p.note(doc, "Cost (Groq): about 1-3 EUR per morning for 10 students. Enable "
                    "pay-as-you-go and set a spend limit on the provider console "
                    "(e.g. https://console.groq.com).")
        p.note(doc, "Upstream resilience: no modification of the OWASP product "
                    "(LiteLLM does the routing). Since PwnzzAI is cloned at a pinned "
                    "commit, this config survives upstream updates, like the rest of "
                    "the sidecar.")


def appendix_docker_en(doc: Document) -> None:
    doc.add_page_break()
    doc.add_heading("Appendix — Install Docker and Docker Compose", level=1)
    p.para(doc, "Do this ONCE per machine, before the install procedure.")

    doc.add_heading("Windows 10/11", level=2)
    p.numbered(doc, "Enable WSL2 (PowerShell as administrator), then reboot:")
    p.code_block(doc, ["wsl --install"])
    p.numbered(doc, "Install Docker Desktop from docker.com (WSL2 backend by default).")
    p.numbered(doc, "Launch Docker Desktop and wait for the 'running' status.")
    p.numbered(doc, "Verify in a terminal:")
    p.code_block(doc, ["docker run --rm hello-world", "docker compose version"])
    p.note(doc, "Virtualization error: enable 'Virtualization / SVM / VT-x' in the "
                "BIOS/UEFI and the Windows feature 'Virtual Machine Platform'. "
                "PowerShell 7+ (pwsh command) is required for the .ps1 scripts; "
                "Windows 10 ships PowerShell 5.1 by default, too old.")

    doc.add_heading("macOS", level=2)
    p.numbered(doc, "Download Docker Desktop for Mac (Apple Silicon or Intel).")
    p.numbered(doc, "Drag Docker into Applications, launch it, wait for 'running'.")
    p.numbered(doc, "Verify:")
    p.code_block(doc, ["docker run --rm hello-world", "docker compose version"])
    p.note(doc, "Apple Silicon (M1/M2/M3): the image builds natively as arm64, no "
                "Rosetta handling required (first build slightly longer).")

    doc.add_heading("Linux (Debian / Ubuntu)", level=2)
    p.para(doc, "Two mutually exclusive methods. Pick one OR the other.")
    p.para(doc, "Method A — distribution packages (recommended for a class):")
    p.code_block(doc, [
        "sudo apt update",
        "sudo apt install -y docker-compose-v2",
        "sudo usermod -aG docker $USER",
        "newgrp docker   # or log out and back in",
        "docker compose version",
    ])
    p.para(doc, "Method B — official Docker repository (latest version):")
    p.code_block(doc, [
        "# set up the repo: https://docs.docker.com/engine/install/ubuntu/",
        "sudo apt install -y docker-ce docker-ce-cli containerd.io \\",
        "    docker-buildx-plugin docker-compose-plugin",
        "sudo usermod -aG docker $USER",
        "newgrp docker",
        "docker compose version",
    ])
    p.note(doc, "The two packages are mutually exclusive. On Ubuntu 25.04+, use "
                "method A (docker-compose-v2). Never mix the two.")
    p.note(doc, "arm64 (Apple Silicon / Snapdragon): if a build fails with "
                "'E: Dynamic MMap ran out of room', it is the APT cache. The "
                "project Dockerfile already includes the fix "
                "(APT::Cache-Start \"100663296\").")


def build_prof_en(tmp: Path) -> None:
    doc = Document()
    p.setup_styles(doc)
    p.add_footer_pagenum(doc, "JuiceLab Coach for PwnzzAI - Teacher guide")
    p.title_block(
        doc, "JuiceLab Coach for PwnzzAI — Teacher guide",
        "PwnzzAI cohort + central JuiceLab dashboard — Windows / macOS / Linux",
    )

    doc.add_heading("1. Your role and the topology", level=1)
    p.para(doc, "Each student runs the full PwnzzAI stack (coach + app + ollama) on "
                "their own machine. While they play, their machines push progress "
                "events to ONE single central dashboard. You get a real-time cohort "
                "matrix (students x labs).")
    p.note(doc, "The dashboard is central and SHARED: a single instance serves both "
                "JuiceLab (Juice Shop) cohorts AND PwnzzAI cohorts — both report into "
                "the same matrix. PwnzzAI NEVER duplicates the dashboard server code; "
                "it plugs in as a client via JUICELAB_DASHBOARD_URL. Anti-pattern to "
                "avoid: deploying a second PwnzzAI-specific dashboard.")
    p.note(doc, "Students NEVER launch the dashboard and never open its interface: "
                "they just point at it with -d <teacher-ip>. Deploying the dashboard "
                "(section 4) is teacher-only.")

    png_archi = tmp / "archi_prof_en.png"
    p.render_mermaid(MERMAID_ARCHI_EN, png_archi, tmp)
    p.add_diagram(doc, png_archi, "Each student machine: coach -> app/ollama, and "
                                  "events reported to the central dashboard.")

    doc.add_heading("2. Prerequisites", level=1)
    section_prerequisites_en(doc)
    p.note(doc, "The dashboard is lightweight. It must be reachable by students: "
                "flat LAN and port " + p.DASHBOARD_PORT + " open in the firewall.")

    doc.add_heading("3. Prepare the cohort", level=1)
    p.para(doc, "Pick a cohort id (JUICELAB_COHORT_ID) that groups students in the "
                "matrix, e.g. " + p.COHORT + ". Then hand each student TWO pieces of "
                "information:")
    p.add_table(
        doc, ["To hand out", "Example", "Role"],
        [
            ["Cohort id", p.COHORT, "Groups the class (students' -c key)"],
            ["Teacher dashboard IP/URL", p.EXAMPLE_IP,
             "Where student machines push their progress (students' -d key)"],
        ],
        [50, 45, 75],
    )
    p.para(doc, "The exact command each student runs (cohort mode):")
    p.code_block(doc, [
        "# Linux / macOS",
        "./scripts/install-student.sh -c " + p.COHORT + " -d " + p.EXAMPLE_IP +
        " -l <firstname>",
        "",
        "# Windows PowerShell 7+",
        ".\\scripts\\install-student.ps1 -Cohort " + p.COHORT +
        " -Dashboard " + p.EXAMPLE_IP + " -Label <firstname>",
    ])
    p.note(doc, "Give each student a unique label (-l / -Label, e.g. their "
                "firstname), otherwise two machines collide in the matrix. The "
                "dashboard port is " + p.DASHBOARD_PORT + " by default; specify "
                "<ip>:<port> if you use another.")

    doc.add_heading("4. Deploy the central dashboard (TEACHER ONLY)", level=1)
    p.note(doc, "This section is teacher-only. Students NEVER run this script: they "
                "have no dashboard to deploy.")
    p.para(doc, "Most common case: a central dashboard already exists (shared "
                "JuiceLab instance). Then you only communicate its IP/URL to students "
                "(section 3). If the dashboard does not exist yet and you want to "
                "deploy it FROM this machine, scripts/deploy-dashboard.sh does a "
                "pinned sparse-checkout of the juicelab repo and pulls ONLY the "
                "teacher part (dashboard/, docker/, scripts/) — the overlay and the "
                "student Juice Shop are never pulled.")
    p.code_block(doc, [
        "# Linux / macOS",
        "./scripts/deploy-dashboard.sh",
        "",
        "# Windows PowerShell 7+",
        ".\\scripts\\deploy-dashboard.ps1",
    ])
    p.para(doc, "The deployment reads these keys from PwnzzAI's .env:")
    p.add_table(
        doc, ["Variable", "Role", "Default"],
        [
            ["JUICELAB_REPO_URL", "juicelab repo (source of the teacher part)",
             p.JUICELAB_REPO_URL],
            ["JUICELAB_DASHBOARD_REF", "pulled ref (SHA advised in prod)", "main"],
        ],
        [60, 65, 65],
    )
    p.para(doc, "Once the dashboard is up, point JUICELAB_DASHBOARD_URL at it "
                "(student-side, that is the value of -d):")
    p.add_table(
        doc, ["Variable (student .env)", "Role", "Example"],
        [
            ["JUICELAB_DASHBOARD_URL", "Dashboard URL as seen from the container",
             "http://host.docker.internal:" + p.DASHBOARD_PORT],
            ["JUICELAB_COHORT_ID", "Cohort id", p.COHORT],
            ["JUICELAB_INSTANCE_LABEL", "Student machine name in the matrix", p.LABEL],
        ],
        [60, 60, 70],
    )
    p.note(doc, "Dashboard on the same machine as the coach: use "
                "host.docker.internal (not 127.0.0.1). Remote dashboard: the "
                "teacher's IP. The dashboard registers the student automatically on "
                "the first event.")
    p.add_screenshot(doc, p.IMG_DIR / "prof-dashboard-light.png",
                     "Teacher dashboard - PwnzzAI cohort on the central instance "
                     "(light theme).")
    p.add_screenshot(doc, p.IMG_DIR / "prof-dashboard-dark.png",
                     "Same matrix in dark theme (toggle via the topbar).")

    doc.add_heading("5. What is reported: the events", level=1)
    p.para(doc, "The coach sends these events via the public POST /api/sync contract "
                "(the same as Juice Shop):")
    p.add_table(
        doc, ["Event", "Trigger", "Data"],
        [
            ["session_start", "opening a lab page", "path, identity if detected"],
            ["hint_revealed", "revealing a hint", "level N1-N5, cost, score"],
            ["challenge_solved", "judge -> Solved", "score, verdict, justification"],
            ["journal_filled", "Journal after save", "length, text"],
            ["quiz_completed", "quiz validation", "score, correct / total"],
            ["badge_earned", "badge unlocked", "badge id"],
        ],
        [45, 60, 65],
    )

    png_seq = tmp / "sequence_en.png"
    p.render_mermaid(MERMAID_SEQUENCE_EN, png_seq, tmp)
    p.add_diagram(doc, png_seq, "From transcript to judge verdict, then reported to "
                                "the dashboard if solved.")

    doc.add_heading("6. Choosing the judge model", level=1)
    p.add_table(
        doc, ["Model", "Verdict", "Recommendation"],
        [
            ["llama3.2:1b", "unreliable (inconsistent)", "avoid for the judge"],
            ["llama3.2:3b", "reliable on clear-cut cases", "advised minimum"],
            ["mistral:7b / llama3.1:8b", "reliable on subtle cases",
             "ideal if the machine can handle it"],
        ],
        [50, 60, 60],
    )
    p.note(doc, "The judge model is independent from the labs' model: labs on 1b "
                "(fast, " + p.LABS_MODEL + ") and judge on 3b (reliable, " +
                p.JUDGE_MODEL + ") on the same machine. Set via COACH_JUDGE_MODEL in "
                ".env.")
    p.note(doc, "The judge and hints always run on local Ollama. The labs' TARGET "
                "(openai_*) can, however, run on a cloud provider (next section).")

    doc.add_heading("7. Configure an LLM provider (target)", level=1)
    section_llm_provider_en(doc, full=True)

    doc.add_heading("8. Teardown", level=1)
    p.para(doc, "Student-side (stopping the PwnzzAI stack on their machine):")
    p.code_block(doc, [
        "./pwnzzai.sh down             # stop (keeps volumes)",
        "./pwnzzai.sh wipe             # down -v: stop + remove volumes",
    ])
    p.para(doc, "Central dashboard deployed from this machine "
                "(in the deployment directory, .juicelab-dashboard/docker):")
    p.code_block(doc, [
        "docker compose --env-file .env down       # stop the dashboard",
        "docker compose --env-file .env down -v    # stop + erase cohort data",
    ])
    p.note(doc, "down -v on the dashboard erases the SQLite database (all cohort "
                "progress). Collect the proofs first if needed.")

    appendix_docker_en(doc)
    doc.save(str(PROF_OUT_EN))
    print(f"  [ok] {PROF_OUT_EN.name} generated ({PROF_OUT_EN.stat().st_size // 1024} Ko)")


def main() -> int:
    print("Generating the EN PwnzzAI TEACHER install guide...")
    with tempfile.TemporaryDirectory(dir=str(DOCS_DIR)) as td:
        build_prof_en(Path(td))
    print("Done.")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
